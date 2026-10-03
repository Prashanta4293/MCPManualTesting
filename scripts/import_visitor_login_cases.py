"""Import only workbook specification columns; never import historical outcomes."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def main():
    source = ROOT/'test_data/Odisha_Tourism_Web_Portal_Test_Cases_v1.0_2026-06-01.xlsx'
    sheet = load_workbook(source,read_only=True,data_only=True)['Test_Cases']
    keys = ['Scenario','Test Case ID','Field Name/Reference','Test Case Description','Preconditions','Steps','Test Data','Expected Result']
    cases, scenario = [], None
    for row_number,row in enumerate(sheet.iter_rows(min_row=4,values_only=True),start=4):
        if row[1]:
            scenario = row[1]
        if str(row[2]).strip() not in {f'TC{i}' for i in range(16,42)}:
            continue
        case = dict(zip(keys,[scenario,*row[2:9]]))
        id = case['Test Case ID']
        case.update({'Module':'Visitor Login','Test Scenario':case['Test Case Description'],'Test Steps':case['Steps'],
                     'Severity':'High' if id in ('TC27','TC32','TC36','TC37','TC39','TC40','TC41') else 'Medium',
                     'automation':'Semi-automated' if id in ('TC36','TC39','TC40') else 'Automated',
                     'negative':id in ('TC27','TC32','TC37','TC40','TC41'),'source':'Test_Cases','source_row':row_number})
        cases.append(case)
    assert len(cases)==26 and len({c['Test Case ID'] for c in cases})==26
    target = ROOT/'data/visitor_login_data.json'
    target.parent.mkdir(exist_ok=True)
    target.write_text(json.dumps({'workbook':source.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'cases':cases},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Imported 26 specification rows; historical Actual Result and Status excluded.')

if __name__=='__main__':
    main()
