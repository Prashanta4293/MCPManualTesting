import json, re
import argparse
from pathlib import Path
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
parser = argparse.ArgumentParser(description='Rebuild the historical manual report from saved evidence, without browser execution.')
parser.add_argument('--output-dir', type=Path, default=ROOT)
OUTPUT = parser.parse_args().output_dir.resolve()
OUTPUT.mkdir(parents=True, exist_ok=True)
E = ROOT / 'odisha_evidence'
BASE = 'https://apptourlfr-stg.estpl.net'
HOME = BASE + '/home'
COLS = ['Test Case ID','Module','Test Scenario','Preconditions','Test Steps','Test Data','Expected Result','Actual Result','Status','Severity','Priority','Tested URL','Evidence/Screenshot','Console Error','Network Error','Remarks']
records = {}
for f in sorted(E.glob('execution_part*.json')):
    for r in json.loads(f.read_text(encoding='utf-8')):
        records[r['index']] = r

# These observations were actually executed on 28 September and are transcribed
# from the Playwright MCP results retained in the conversation. Screenshots exist.
extra = [
 (48,'International Sand Art Festival 2026','international-sand-art-festival','Plan Your Trip','Highlighted Events','Grains of Sea Shaped into Breathtaking Art','Festival summary, venue Chandrabhaga Beach, dates 1st to 5th December, gallery and visit information.','Swiper is not defined'),
 (49,'Odisha Parab','odisha-parab','Plan Your Trip','Highlighted Events',"Where Odisha’s soul comes alive beyond its borders",'Welcome to Odisha Parab, Varanasi; event dates 22–24 August 2026 and Varanasi venue. About section says “The Ahmedabad edition”.',''),
 (50,'Mahaprasad','mahaprasad','Plan Your Trip','Cuisines of Odisha','Mahaprasad','Overview, heritage, ingredients, preparation, serving style and where to try. Labels a rice/lentil/vegetable sacred meal “LEGENDARY SWEET”.','Failed to load resource: 404 patter-1.png'),
 (51,'Rasagola','rasagola','Plan Your Trip','Cuisines of Odisha','Rasagola','Overview, cultural significance, recipe, ingredients, preparation and where to try.','Failed to load resource: 404 patter-1.png'),
 (52,'Chhenapoda','chhenapoda','Plan Your Trip','Cuisines of Odisha','Chhenapoda','Overview, baked cheese dessert, cultural significance, recipe and where to try.','Failed to load resource: 404 patter-1.png'),
 (53,'Sarsatia','sarsatia','Plan Your Trip','Cuisines of Odisha','Sarsatia','Overview, Sambalpur sweet, traditional recipe, serving style and where to try.','Failed to load resource: 404 patter-1.png'),
 (54,'About Us','about-us','Hamburger','About Odisha Tourism','About Us','About Department and Organizational Structure.',''),
 (56,'Vision, Mission & Objectives','vision-mission-objectives','Hamburger','About Odisha Tourism','Vision, Mission & Objectives','Vision, mission, objectives, strategic pillars and key objectives.',''),
 (57,'Odisha Tourism Policy','odisha-tourism-policy','Hamburger','About Odisha Tourism','Odisha Tourism Policy','Five documents including Handbook on Odisha Tourism Policy 2022 (Amendments) 2026 with dates and Download links.',''),
 (58,'Operational Guidelines','operational-guidelines','Hamburger','About Odisha Tourism','Operational Guidelines','Eleven guidelines with dates and PDF download links.',''),
 (59,'Tenders and Notices','tenders-and-notices','Hamburger','About Odisha Tourism','Tenders & Notices','Category, keyword/date filters, populated tender cards and document links.',''),
 (60,'Newsletter','newsletter','Hamburger','Media & Updates','Newsletter','Odisha Unravelled editions April 2026, January 2026 and January 2025.',''),
 (61,'Brochure','brochure','Hamburger','Media & Updates','Brochures','Explore Brochures: Beaches, Buddhist Circuit, Eco Retreats and other PDF brochure cards.',''),
 (62,'Photo Gallery','photo-gallery','Hamburger','Media & Updates','Photo Gallery','The Odisha Collection, category filters and populated image gallery.',''),
 (63,'Video Gallery','video-gallery','Hamburger','Media & Updates','Video Gallery','Odisha Through the Lens and visible video thumbnails; visually reviewed screenshot.',''),
 (64,'Traveller Blogs','blogs-stories','Hamburger','Media & Updates','Traveller Blogs','Populated travel blog cards, authors, dates, excerpts and Trending Now.',''),
 (65,'Festival & Events','events','Hamburger','Media & Updates','Events & Festivals','Event calendar, top ten events and frequently asked questions.',''),
 (66,'Privacy Policy','privacy-policy','Hamburger','Policies & Legal','Privacy Policy','Policy sections including data collection, retention, cookies and contact information. Legal correctness not assessed.',''),
 (67,'Disclaimer','disclaimer','Hamburger','Policies & Legal','Disclaimer','Website development, accuracy, third-party information, liability and contact sections.',''),
 (68,'Contact Us','contact-us','Hamburger','Help & Support','Contact Us','Department of Tourism address, phone, email, social links and map. No contact submission.',''),
 (69,'Site Map','site-map','Hamburger','Help & Support','Sitemap','Hierarchical portal directory with Discover, Experience and Plan Your Trip links.',''),
 (70,'FAQs','faqs','Hamburger','Help & Support','Frequently Asked Questions','Ten visible questions and pagination; answer-expansion behavior not part of destination-load check.',''),
]
for i,name,slug,menu,group,heading,content,err in extra:
    records[i] = dict(index=i,name=name,href=BASE+'/'+slug,menu=menu,group=group,console=[err] if err else [],network=[dict(url=BASE+'/patter-1.png',status=404)] if 'patter' in err else [],documents=[],state=dict(url=BASE+'/'+slug,title='Recorded in MCP transcript; see page heading',headings=[heading],breadcrumbs=[group+' / '+name],content=content,broken=[],pendingImages=[]),screenshot=f'manual_test_cases/odisha_evidence/nav-{i:02}.png',transcribed=True)

rows=[]
def add(id,module,scenario,steps,expected,actual,status='Passed',severity='N/A',priority='P2',url=HOME,evidence='',console='None observed',network='None observed',remarks='',data='None',pre='Anonymous user; Playwright MCP browser; home page available'):
    rows.append(dict(zip(COLS,[id,module,scenario,pre,steps,data,expected,actual,status,severity,priority,url,evidence,console,network,remarks])))

def shorturl(s):
    return s.split('?')[0] if 'google' in s or 'fbcdn' in s else s

nav_defects={
 1:('D02','Content contains inconsistent coastline lengths: nearly 480 km in Overview versus nearly 574 km in Geography.','Low','P3'),
 3:('D03','Public content includes an editorial instruction: “The image used is of the Hathigumpha, Udaygiri Caves and the information displayed should reflect the above historical context in the text portion.”','Low','P3'),
 18:('D04','Central Odisha reuses the heading “DISTRICTS THAT SHAPE THE COAST” despite its inland content.','Low','P3'),
 20:('D04','Western Odisha reuses the heading “DISTRICTS THAT SHAPE THE COAST” despite its inland content.','Low','P3'),
 29:('D05','Map image viewer renders; console reports “Map elements not found”. No page heading or breadcrumb in viewer; not a blank page.','Medium','P2'),
 34:('D06, D07','Console: Swiper is not defined; YouTube maxresdefault thumbnail returns HTTP 404.','Medium','P2'),
 36:('D08','Tour Operators directory is incorrectly labelled “Government Approved Tour Guides” with a guide-specific introduction.','Low','P3'),
 44:('D06','Console: Swiper is not defined.','Medium','P2'),
 46:('D06','Console: Swiper is not defined.','Medium','P2'),
 47:('D06','Console: Swiper is not defined.','Medium','P2'),
 48:('D06','Console: Swiper is not defined. Destination verified after enlarging viewport; see separate menu overflow test.','Medium','P2'),
 49:('D09','Same event page identifies Varanasi in title/venue and Ahmedabad in About text.','Medium','P2'),
 50:('D10, D11','Missing patter-1.png (404); Mahaprasad is labelled “LEGENDARY SWEET” despite meal description.','Low','P3'),
 51:('D10','Missing patter-1.png (404).','Low','P3'),
 52:('D10','Missing patter-1.png (404).','Low','P3'),
 53:('D10','Missing patter-1.png (404).','Low','P3')}

for i,r in sorted(records.items()):
    s=r['state']; defect=nav_defects.get(i)
    heads='; '.join(s.get('headings',[])[:7])
    crumbs='; '.join(dict.fromkeys(s.get('breadcrumbs',[]))).replace('\n',' ')
    content=re.sub(r'\s+',' ',s.get('content',''))[:1700]
    if i==29: content='A large Odisha tourist map image, close icon and download icon are visible in the screenshot.'
    actual=f"Menu: {r['menu']}; Group: {r.get('group','N/A')}; Submenu: {r['name']}. Final URL: {s['url']}. Heading(s): {heads or 'None (map viewer)'}. Breadcrumb: {crumbs or 'None (map viewer)'}. Content: {content}. No blank page or page-level 404/500 observed. No unexpected redirect. Loaded visible IMG failures: {len(s.get('broken',[]))}."
    if defect: actual+=' Finding: '+defect[1]
    network='\n'.join(f"{n.get('status',n.get('error',''))}: {shorturl(n['url'])}" for n in r.get('network',[])) or 'None observed'
    remarks='28 Sep 2026. Expected URL derived from clicked link href; no separate requirements baseline supplied. Return to home executed before next item. Destination smoke check only; inner-page forms, downloads, pagination and media playback not exercised.'
    if r.get('transcribed'): remarks+=' Detailed result transcribed from executed MCP conversation; screenshot retained.'
    if any('analytics.google' in n.get('url','') for n in r.get('network',[])): remarks+=' Analytics ERR_ABORTED recorded separately; not treated as a functional destination failure.'
    if i in [13,14,15,53]: remarks+=' Initial animation-related click timeout resolved on retry.'
    if i in [1,2,3,4]: remarks+=' Initial offscreen lazy images pending; 29 Sep supplemental scroll audit confirmed all inspected IMG resources loaded (image-rechecks.json).'
    add(f'NAV-{i:03}',r['menu']+(' / '+r['group'] if r.get('group') else ''),'Open '+r['name'],'Open home; open/hover menu; click named item; inspect destination heading, breadcrumb, content and images; collect console/network; screenshot; return home.',f"Reach {r['href']} with relevant content, no page-level error, unexpected redirect, broken loaded images or runtime/resource errors.",actual,'Failed' if defect else 'Passed',defect[2] if defect else 'N/A',defect[3] if defect else 'P2',s['url'],r['screenshot'], '\n'.join(r.get('console',[])) or 'None observed',network,remarks+((' Defect '+defect[0]+'.') if defect else ''),r['href'])

externals=[
 (55,'Department of Tourism, Government of Odisha','https://dot.odisha.gov.in/en/','Home | Department of Tourism','Department of Tourism / Government of Odisha; populated official department content.','Failed','404 resources; canlendar.js MIME error; Facebook widget errors; certificate failures in embedded media.','HTTP 404 for canlendar.js and another resource; embedded fbcdn requests ERR_CERT_AUTHORITY_INVALID.','EXT02; external site ownership.'),
 (71,'Book Now','https://www.bookodisha.com/','Odisha Tourism | Home','Discover the Best Kept Secrets; hotel, package and ticketing content.','Failed','Clarity script blocked by Content Security Policy.','Clarity CSP failure; Google tag ERR_BLOCKED_BY_ORB; analytics abort.','EXT01; correct destination opens, external diagnostic defect; no booking form submitted.'),
 (72,'Facebook','https://www.facebook.com/OdishaTourismOfficial','Odisha Tourism | Facebook','Odisha Tourism official page and intro visible.','Failed','ErrorUtils: Action bar is null on Comet.','None observed','EXT03; Facebook share URL resolves to official profile as expected; anonymous login prompt; no login.'),
 (73,'X','https://x.com/odisha_tourism?s=11&t=4pIKcmQIe3N0gnvIhlPZjA','Privacy error','Chrome privacy interstitial: Your connection is not private; NET::ERR_CERT_AUTHORITY_INVALID.','Blocked','None captured; browser certificate interstitial.','NET::ERR_CERT_AUTHORITY_INVALID displayed by browser.','External certificate/environment blocker. Did not bypass warning. Final browser URL chrome-error://chromewebdata/; profile content and images unverified.'),
 (74,'Instagram','https://www.instagram.com/odishatourismofficial?igsh=NHBubHhoMjl0d2lj','Odisha Tourism (@odishatourismofficial) • Instagram photos and videos','Official profile, description and profile content visible.','Passed','None observed','None observed','Cookie banner intercepted initial click; dismissed and retried successfully. No login.'),
 (75,'YouTube','https://www.youtube.com/@odishatourismofficial','Odisha Tourism - YouTube','Odisha Tourism channel with Videos, Popular videos and Shorts.','Failed','Resource returned HTTP 401.','Passive Google sign-in request returned HTTP 401; youtube.com request aborted.','EXT04; expected channel canonicalization from youtube.com to www.youtube.com. Empty-src lazy image placeholders not classified as broken resources. No sign-in or subscription.')]
for i,name,url,title,content,status,console,network,remark in externals:
    add(f'NAV-{i:03}','Header / External',name,'Home; open header/hamburger; click link; verify external confirmation destination; Continue; inspect new tab; capture screenshot and diagnostics; close tab; return home.','Expected official external destination opens with relevant content and no browser/resource errors.',f'Title: {title}. {content} Breadcrumb: N/A for external home/profile. No loaded nonempty-src broken IMG elements detected during observation.' if i!=73 else content,status,'Medium' if status=='Failed' else 'N/A','P2',url,f'manual_test_cases/odisha_evidence/external-{i}.png',console,network,'28 Sep 2026. '+remark+' Popup diagnostics attached after popup creation; earliest navigation response may not be captured.',url)

add('HOME-001','Home','Home URL, title and visible sections','Navigate to home; take snapshot; inspect URL/title; scroll main sections.','Home loads with Odisha Tourism branding and relevant main content.','Home loads at /home with title “Odisha Tourism: Official Website to Plan Your Travel & Holiday”. Experience Odisha, Uniquely Odisha, Odisha Walks, Choose your vibe, Navigate Odisha, Glimpses of Odisha, Video Gallery, Postcards and Travel Diaries observed.',evidence='manual_test_cases/odisha_evidence/home-final.png',remarks='28 Sep initial execution; 29 Sep home-content follow-up. No page-level 404/500 or unexpected redirect. Hero text changes with animation.')
add('HOME-002','Home','Favicon asset','Read icon link; fetch referenced icon in browser; decode image.','Favicon is configured and resolves to a decodable image.','Both icon links reference odisha-tourism-favicon.png; HTTP 200 image/png; decoded size 136×146.',evidence='manual_test_cases/odisha_evidence/home.png',remarks='28 Sep. Browser tab chrome itself is outside page screenshot; configured resource and decoding verified, not tab-level visual rendering.')
add('HOME-003','Home','Home image and network health','Scroll page; inspect loaded IMG dimensions; inspect console and network failures.','Visible images load and home resources complete successfully.','No loaded visible IMG failures in 29 Sep audit. Hidden Video Gallery Background request to loremflickr.com repeatedly fails with ERR_BLOCKED_BY_ORB. Hidden empty Event Image placeholder is not counted as a visible defect. Some offscreen carousel images remained lazy/pending.','Failed','Low','P3',evidence='manual_test_cases/odisha_evidence/home-final.png',network='https://loremflickr.com/600/400/travel,nature?... => net::ERR_BLOCKED_BY_ORB',remarks='D01. No visible gallery outage proven; video background still renders. See home-final-network.txt. Does not assert every carousel/background image was loaded.')
add('HDR-001','Header','Logo image and navigation','Click header logo; inspect URL, title and loaded image.','Logo loads and returns to home.','Logo decoded (natural width 442); click reaches /web/guest/home with the same home title and content.',url=BASE+'/web/guest/home',evidence='manual_test_cases/odisha_evidence/logo.png',remarks='28 Sep. /web/guest/home is the advertised home alias, not an unexpected redirect.')
add('HDR-002','Header','Discover menu reveal','Hover Discover from home; inspect submenu; click its items in NAV-001–008.','All eight Discover links are visible and actionable.','Eight links revealed; each clicked and destination assessed in navigation rows.',evidence='manual_test_cases/odisha_evidence/nav-01.png',remarks='28 Sep. Menu reveal only; destination defects counted in individual rows.')
add('HDR-003','Header','Experience menu reveal','Hover Experience; wait for animation; inspect eight items.','Eight submenu links available.','Eight items clicked and assessed in NAV-009–016.',evidence='manual_test_cases/odisha_evidence/nav-13.png',remarks='28 Sep. Three initial animation timeouts resolved on retry.')
add('HDR-004','Header','Plan Your Trip menu fits viewport','At 1366×577 hover Plan Your Trip; attempt Sand Art Festival link; inspect menu geometry; repeat destination at 1366×900.','Menu items remain reachable by pointer or menu scrolling.','Menu bottom is 680px in a 577px viewport; overflow is visible with no internal scrolling. Sand Art link starts at y=582.9 and click times out outside viewport. Destination subsequently reached at 1366×900.','Failed','Medium','P2',evidence='manual_test_cases/odisha_evidence/menu-overflow.png',remarks='D12. 28 Sep; distinguishes menu failure from destination result.',data='1366×577; recovery at 1366×900')
add('HDR-005','Header','Hamburger opens and closes','Click hamburger; inspect groups/links; click Close.','Menu opens and closes without navigation.','About Odisha Tourism, Media & Updates, Policies & Legal and Help & Support appear; Close hides menu. All 17 visible menu links and four social links tested individually.',evidence='manual_test_cases/odisha_evidence/cookie-overlap.png',remarks='28 Sep. Hidden mobile duplicate Discover/Experience/Plan menus were not counted as desktop-visible links.')
for code,name,status in [('Hi','Hindi','Failed'),('Od','Odia','Failed'),('En','English','Passed')]:
    add('LANG-'+code,'Header / Language','Select '+name,'Return home; open language selector; select option; wait; inspect label, URL, html language and visible headings.','Selected language is applied to label and visible page content.',f'Selected {name}; URL remains /home, label EN, document language en-US, heading Experience Odisha and English content remain.',status,'Medium' if status=='Failed' else 'N/A','P2',evidence=f'manual_test_cases/odisha_evidence/language-{code}.png',remarks='28 Sep. '+('D13. Language selection has no visible effect.' if status=='Failed' else 'Selecting already-active English verified; no claim of successful round-trip translation.'),data=name)
add('SEARCH-001','Header / Search','Open search overlay without runtime error','Click header search icon; inspect overlay and console.','Search opens with input and no runtime exception.','Overlay opens but onclick throws TypeError: Cannot read properties of null (reading value).','Failed','Medium','P2',evidence='manual_test_cases/odisha_evidence/search-console-error.png',console="TypeError: Cannot read properties of null (reading 'value') at HTMLButtonElement.onclick (home:3427:328)",remarks='D14. 28 Sep. Search remains usable despite error.')
add('SEARCH-002','Header / Search','Search Puri','Open search; enter Puri; click Search; inspect result URL, heading, breadcrumb and result list.','Search query is preserved and relevant results displayed.','Reached /search-result?q=Puri; title Search Result; heading Search; Home / Search Results breadcrumb; 15 results including Puri, Puri beaches and Shree Jagannath Temple.',url=BASE+'/search-result?q=Puri',evidence='manual_test_cases/odisha_evidence/search-puri.png',remarks='28 Sep. No error console messages on result page. Search snippets contain formatting artifacts; exact ranking and every result link not verified.',data='Puri')
add('ACCOUNT-001','Header / Profile','Login dialog display and close','Click profile icon; inspect form; close without input or submit.','Anonymous login dialog opens and closes.','Login heading, email/mobile, password, captcha, Login, Forgot Password, Sign Up and Official Login visible. Close hides dialog.',evidence='manual_test_cases/odisha_evidence/login-dialog.png',remarks='28 Sep. No credentials entered; no form submission.')
for code,name,path,text in [('ForgotPassword','Forgot Password?', '/home','Forgot Password view with registered email/OTP instruction and Back to Login; no OTP requested.'),('SignUp','Sign Up','/web/guest/visitor-registration','Visitor Registration heading, breadcrumb, Register Here form, required fields and guidelines visible.'),('OfficialLogin','Official Login','/web/guest/administrator-login','Official Portal Login, DTO User/State Department/Administrator choices and credential/captcha controls visible.')]:
    add('ACCOUNT-'+code,'Header / Profile',name+' opens','Return home; open profile; click '+name+'; inspect resulting UI, URL, headings/breadcrumb, loaded images and runtime/network diagnostics; return home.','Correct account view opens without submitting a form.',text+' Follow-up found no loaded visible IMG failures, console errors or failed requests. No unexpected redirect. '+('Registration breadcrumb: Visitor Registration.' if code=='SignUp' else 'Breadcrumb N/A for modal or standalone login; Official Portal Login uses visible text rather than h1/h2/h3.'),url=BASE+path,evidence=f'manual_test_cases/odisha_evidence/account-final-{code}.png',remarks='28 Sep initial UI check; 29 Sep diagnostic follow-up saved in account-final.json. No credentials entered or OTP/form submitted.')
add('ACC-001','Header / Accessibility','Accessibility panel opens and closes','Click Accessibility Options; inspect panel; click Close Accessibility Panel.','Options panel is displayed and dismissible.','Panel displays text sizing, contrast, invert, spacing, image/link/cursor, screen reader and reset controls; closes.',evidence='manual_test_cases/odisha_evidence/accessibility.png',remarks='28 Sep. Option effects and assistive-technology compatibility are not implied by this panel display check.')
add('TOP-001','Header / Announcements','Collapse and restore announcement bar','Click Hide Topbar; inspect collapsed state; click Show Announcements.','Announcement bar can be hidden and restored.','Bar moves above viewport (button y=-30.5); restore tab at y=0; clicking it returns toggle button to y=7.5.',evidence='manual_test_cases/odisha_evidence/topbar-collapsed.png',remarks='29 Sep. Collapse state persists across reload; initial Events click while collapsed retried after restoring. No site defect.')
add('TOP-002','Header / Announcements','Events Of Odisha destination','Restore bar; return home; click Events Of Odisha; inspect destination and diagnostics; return home.','Advertised /web/guest/events opens a populated events calendar.','Reached /web/guest/events; title Odisha Tourism: Events & Festivals Calendar; Events & Festivals and Events Calendar headings. Breadcrumb: Media & Updates / Festival & Events. No loaded IMG failures, console errors or request failures in observation.',url=BASE+'/web/guest/events',evidence='manual_test_cases/odisha_evidence/topbar-events.png',remarks='29 Sep. Expected home-site alias, no unexpected redirect. Breadcrumb verified with broader selector after .breadcrumb alone returned none.')
add('A11Y-001','Header / Keyboard','Skip to Main Content focus transfer','Focus skip link; press Enter; wait; inspect focus; press Tab.','Focus skips header navigation and proceeds into main content.','After Enter focus remains on skip anchor, URL has no hash and scrollY=0. Next Tab focuses header logo; header navigation was not skipped.','Failed','Medium','P2',evidence='manual_test_cases/odisha_evidence/skip-link-focus.png',remarks='D15. Reproduced 29 Sep. Main-content target exists but keyboard focus transfer fails.')
add('LIMIT-001','Scope limitations','Prohibited form submissions','Do not submit registration, login, booking, payment, contact, feedback or OTP forms.','No submissions in this execution.','Not executed by explicit user constraint.','Not Tested',priority='N/A',console='Not tested',network='Not tested',remarks='Deliberate exclusion; UI display/navigation may be tested separately.')
add('LIMIT-002','Scope limitations','Deep accessibility, mobile and media coverage','Not part of desktop header navigation smoke execution.','Separate tests required for option effects, screen-reader speech, mobile menus and exhaustive carousel/media states.','No mobile breakpoint run; no native screen reader/audio assessment; no exhaustive background-image, offscreen carousel image or video playback validation.','Not Tested',priority='N/A',console='Not tested',network='Not tested',remarks='Playwright page screenshots do not expose browser favicon chrome or audible output. Human/assistive technology review needed for those aspects.')
add('SEARCH-003','Header / Search','Close search overlay','Open search; click Close; wait for animation; inspect input visibility and bounds.','Overlay closes and home remains displayed.','Search input becomes invisible with 0×0 bounds; URL stays /home.',evidence='manual_test_cases/odisha_evidence/search-closed.png',console='Search opening TypeError already recorded in SEARCH-001',remarks='29 Sep. Close behavior passed independently of opener defect. Immediate visibility during closing animation was not treated as failure.')
add('HDR-006','Header / External','Cancel external navigation','Click Book Now; inspect external confirmation; choose Cancel.','Dialog closes without opening external tab.','Modal count becomes zero; one browser tab remains and URL stays /home.',evidence='manual_test_cases/odisha_evidence/account-final.json',remarks='29 Sep. No navigation or form submission triggered by Cancel.')

defect_data=[
 ('D01','Hidden home background resource blocked','Low','P3','HOME-003','Remove/repair loremflickr background request; visible video background still works.','home-final.png'),
 ('D02','Conflicting coastline lengths on same page','Low','P3','NAV-001','Overview says nearly 480 km; geography section says nearly 574 km. Confirm authoritative content.','nav-01.png'),
 ('D03','Editorial instruction exposed in public content','Low','P3','NAV-003','Remove author instruction beneath Ancient & Medieval Monuments.','nav-03.png'),
 ('D04','Coastal template heading on inland regional pages','Low','P3','NAV-018, NAV-020','Central and West Odisha use DISTRICTS THAT SHAPE THE COAST. Content review needed.','nav-18.png'),
 ('D05','Tourist Map logs missing-elements error','Medium','P2','NAV-029','Map renders as image viewer. Console error exists; absence of heading/breadcrumb may be intentional viewer design, not a blank-page failure.','nav-29.png'),
 ('D06','Swiper initialization error on five destinations','Medium','P2','NAV-034, NAV-044, NAV-046, NAV-047, NAV-048','Swiper is not defined. Carousel functional impact was not separately tested.','nav-34.png'),
 ('D07','Heritage Properties video thumbnail HTTP 404','Low','P3','NAV-034','https://img.youtube.com/vi/-aPINUood2M/maxresdefault.jpg fails; fallback visual impact not established.','nav-34.png'),
 ('D08','Tour Operators content labelled Tour Guides','Low','P3','NAV-036','Heading and introduction refer to guides while table lists tour operators/travel agents.','nav-36.png'),
 ('D09','Odisha Parab mixes Varanasi and Ahmedabad','Medium','P2','NAV-049','Venue/intro say Varanasi; About text says Ahmedabad edition.','nav-49.png'),
 ('D10','Four cuisine pages request missing pattern image','Low','P3','NAV-050, NAV-051, NAV-052, NAV-053','/patter-1.png returns 404 on Mahaprasad, Rasagola, Chhenapoda and Sarsatia.','nav-50.png'),
 ('D11','Mahaprasad has sweet-category label','Low','P3','NAV-050','LEGENDARY SWEET conflicts with page description of a sacred meal of rice, lentils and vegetables.','nav-50.png'),
 ('D12','Plan Your Trip menu overflows short desktop viewport','Medium','P2','HDR-004','At 1366×577 bottom links fall below viewport; no internal menu scrolling; click times out.','menu-overflow.png'),
 ('D13','Hindi and Odia selection has no visible effect','Medium','P2','LANG-Hi, LANG-Od','Both leave EN, en-US and English content unchanged.','language-Hi.png'),
 ('D14','Search opener throws null-value TypeError','Medium','P2','SEARCH-001','Search overlay opens but inline onclick accesses missing searchKeyword element.','search-console-error.png'),
 ('D15','Skip link does not bypass header keyboard focus','Medium','P2','A11Y-001','After activation the next Tab goes to header logo, not main content.','skip-link-focus.png'),
 ('EXT01','Book Odisha third-party script failures','Low','P3','NAV-071','External ownership: Clarity CSP and Google tag failures; booking home still renders.','external-71.png'),
 ('EXT02','Department site runtime/resource failures','Medium','P2','NAV-055','External ownership: 404 script/MIME errors and embedded media certificate errors.','external-55.png'),
 ('EXT03','Facebook runtime error','Low','P3','NAV-072','External ownership: Action bar is null on Comet; official profile visible.','external-72.png'),
 ('EXT04','YouTube passive sign-in request unauthorized','Low','P3','NAV-075','External ownership: HTTP 401 in passive sign-in flow; intended channel content visible.','external-75.png')]
counts=Counter(r['Status'] for r in rows)

wb=Workbook(); summary=wb.active; summary.title='Execution Summary'
for row in [['Odisha homepage exploratory execution','28–29 September 2026'],['Target',HOME],['Method','Playwright MCP, anonymous Chrome browser; initial 1366×577 viewport, later 1366×900'],['Total',len(rows)],*[ [s,counts[s]] for s in ['Passed','Failed','Blocked','Not Tested']],['Primary navigation items',75],['Staging defect groups',15],['External diagnostic groups',4],['Blocked destination','X certificate interstitial; not bypassed'],['Evidence','odisha_evidence/'],['Scope','Desktop home/header exploration; no prohibited form submissions'],['Status rule','Passed only for stated executed scope. Failed if observed content/runtime/resource defect. Analytics abort alone not a functional failure.'],['Dates','Initial primary navigation: 28 Sep. Home, image and topbar/keyboard follow-up: 29 Sep. Not a full rerun on 29 Sep.']]: summary.append(row)
ws=wb.create_sheet('Test Execution');ws.append(COLS)
for r in rows: ws.append([r[c] for c in COLS])
ds=wb.create_sheet('Defect Summary');ds.append(['Defect ID','Title','Severity','Priority','Affected Test Cases','Finding / qualification','Evidence/Screenshot'])
for d in defect_data: ds.append(list(d[:-1])+['odisha_evidence/'+d[-1]])
ls=wb.create_sheet('Limitations');ls.append(['Area','Execution boundary'])
for x in [
 ('Primary navigation','All 53 desktop mega-menu links, 17 hamburger links, Book Now and 4 hamburger social links were clicked. One X destination is blocked.'),
 ('Expected URL','Derived from rendered link href. No approved design/requirements supplied; matching route/content is the practical oracle.'),
 ('External redirects','Facebook share-to-profile and YouTube host canonicalization accepted. X certificate interstitial prevents content verification.'),
 ('Favicon','Icon link, HTTP 200 and PNG decoding verified. Native tab-chrome favicon visual cannot be captured by page screenshot.'),
 ('Images','Visible loaded IMG natural dimensions and network failures checked. Hidden placeholders, pending carousel images and every CSS background are not exhaustively verified.'),
 ('Content','Relevant headings/breadcrumbs/text verified, not historical/legal/medical/travel factual accuracy.'),
 ('Account views','UI and scoped console/network/image audit performed on 29 Sep for Forgot Password, Sign Up and Official Login. No submissions.'),
 ('Submissions','No registration/login/booking/payment/contact/feedback/OTP submissions.'),
 ('Accessibility','Panel display and keyboard skip tested. Option effects, speech and native assistive technology not fully verified.'),
 ('Viewport','Desktop only; no mobile breakpoint test. Short viewport overflow documented before enlarging.'),
 ('Evidence provenance','First 47 detailed records saved as JSON. Later executed observations transcribed from prior MCP outputs with retained screenshots. Recheck records dated 29 Sep.'),
 ('Diagnostics','External popup listeners attached after popup creation: earliest navigation response may be missed. Internal analytics request aborts recorded, not classified as page failures.'),
 ('Map viewer','Text extraction initially looked empty; screenshot proves rendered map. Blank-page claim explicitly corrected.'),
 ('Date separation','28 Sep execution results remain historical observations. Only identified follow-up tests rerun 29 Sep.')]: ls.append(x)

fills={'Passed':'E2F0D9','Failed':'FCE4D6','Blocked':'FFF2CC','Not Tested':'E7E6E6'}
for sheet in wb:
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    for c in sheet[1]: c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value,str): c.value=re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]','',c.value)[:32760]
            c.alignment=Alignment(vertical='top',wrap_text=True)
    sheet.sheet_view.zoomScale=80
    for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=28
summary.column_dimensions['A'].width=30;summary.column_dimensions['B'].width=110
widths=[18,30,38,42,62,42,65,110,16,16,12,65,58,70,70,95]
for n,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(n)].width=w
for row in ws.iter_rows(min_row=2):
    row[8].fill=PatternFill('solid',fgColor=fills[row[8].value]); row[8].font=Font(bold=True)
    ws.row_dimensions[row[0].row].height=125
    if row[12].value:
        rel=row[12].value.replace('manual_test_cases/','');row[12].hyperlink=rel;row[12].font=Font(color='0563C1',underline='single')
    row[11].hyperlink=row[11].value;row[11].font=Font(color='0563C1',underline='single')
ds.column_dimensions['B'].width=55;ds.column_dimensions['E'].width=52;ds.column_dimensions['F'].width=110;ds.column_dimensions['G'].width=55
for row in ds.iter_rows(min_row=2):
    row[6].hyperlink=row[6].value;row[6].font=Font(color='0563C1',underline='single');ds.row_dimensions[row[0].row].height=70
ls.column_dimensions['B'].width=135
for sheet,name in [(ws,'OdishaTests'),(ds,'OdishaDefects')]:
    table=Table(displayName=name,ref=sheet.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);sheet.add_table(table)
dest=OUTPUT/'odisha_homepage_execution.xlsx';wb.save(dest)

md=['# Odisha homepage exploratory execution','',f'Executed **28–29 September 2026** using Playwright MCP against {HOME}.','',f"**Total: {len(rows)} | Passed: {counts['Passed']} | Failed: {counts['Failed']} | Blocked: {counts['Blocked']} | Not tested: {counts['Not Tested']}**",'',
 '**Coverage:** 75 primary header navigation items: 53 mega-menu destinations, 17 hamburger links, Book Now, and four social links. Returned to home between items. Home branding/title/favicon asset, logo, language, search, profile views, hamburger, accessibility panel, announcement controls and keyboard skip were also assessed. No registration, login, booking, payment, contact, feedback or OTP form was submitted.','',
 'Primary navigation ran on 28 September. Follow-up home, image, announcement and keyboard checks ran on 29 September; this is not a full rerun of every destination on the second day. Desktop viewport started at 1366×577 and was enlarged to 1366×900 to complete links clipped by the mega-menu.','',
 'The home page loads. No tested internal destination was a confirmed blank page or page-level 404/500. Tourist Map is an image viewer, not blank. X content verification was blocked by a certificate interstitial. Some otherwise reachable pages failed their combined test because of content, JavaScript or resource errors.','',
 '## Defect summary','', '| ID | Finding | Severity / Priority | Tests | Evidence |','|---|---|---|---|---|']
for did,title,sev,pri,tests,note,file in defect_data:
    md.append(f'| {did} | {title} — {note} | {sev} / {pri} | {tests} | [Screenshot](odisha_evidence/{file}) |')
md += ['', '15 staging-site defect groups and 4 external-site diagnostic groups. These are distinct from the number of failed test cases; one shared defect can affect multiple cases. Severity and priority are provisional QA assessments.', '',
 '## Blocked and not verified','',
 '- **Blocked:** X opens Chrome’s `NET::ERR_CERT_AUTHORITY_INVALID` privacy warning. No bypass; profile content, images and normal URL completion are unverified.',
 '- **Not tested by instruction:** form submissions and authentication/booking/payment completion.',
 '- **Not tested:** mobile menus, exhaustive offscreen/carousel/CSS-background images, media playback, accessibility option effects and native screen-reader audio.',
 '- **Favicon limitation:** configured PNG was fetched with HTTP 200 and decoded at 136×146; actual browser-tab chrome appearance is not visible to Playwright page screenshots.',
 '- Supplemental account views and event-alias breadcrumb were checked on 29 September. No authenticated profile state or form completion was verified.',
 '- No approved specification was supplied. Destination URLs were compared to clicked link targets; factual/legal accuracy, inner-page forms, filters, pagination and downloadable documents were outside this header smoke execution.',
 '- Analytics `ERR_ABORTED` entries alone were not counted as functional failures. External popup diagnostics can miss the earliest response because listeners were installed after the popup appeared.',
 '- Some lazy carousel images remain pending outside the viewed state. No blanket claim that every image/media state passed is made.',
 '', '## Test results','', '| ID | Module | Scenario | Status |', '|---|---|---|---|']
for r in rows: md.append('| '+' | '.join(str(r[c]).replace('|','/').replace('\n',' ') for c in ['Test Case ID','Module','Test Scenario','Status'])+' |')
md += ['', '## Artifacts','', '- [Excel execution report](odisha_homepage_execution.xlsx): all 16 requested columns, summary, defects and limitations.', '- [Evidence directory](odisha_evidence/): screenshots, detailed JSON observations and home console/network logs.', '', 'Passed means the assertions stated in that row were actually executed and observed. It does not imply untested destination features or external integrations were verified.']
(OUTPUT/'odisha_homepage_summary.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
(OUTPUT/'odisha_evidence').mkdir(exist_ok=True)
(OUTPUT/'odisha_evidence'/'consolidated_test_rows.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
check=load_workbook(dest,read_only=True)
assert list(next(check['Test Execution'].values))==COLS
assert check['Test Execution'].max_row==len(rows)+1
missing=[r['Evidence/Screenshot'] for r in rows if r['Evidence/Screenshot'] and not (ROOT.parent/r['Evidence/Screenshot']).exists()]
assert not missing,missing
assert len({r['Test Case ID'] for r in rows})==len(rows)
print(json.dumps({'total':len(rows),'counts':dict(counts),'defect_groups':len(defect_data),'primary_navigation':len([r for r in rows if r['Test Case ID'].startswith('NAV-')]),'report':str(dest),'missing_evidence':missing},indent=2))
